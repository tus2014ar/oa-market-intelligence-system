"""Loader for the Branded Generic product reference workbooks (OA and RA).

Each is a flat table with one row per (Manufacturer, Product, Brand/Generic tag) and one
visit-count column per ICD-10 code (docs/data_dictionary.md §4). It is the product
master: it says which products exist and how IQVIA tags them, not how they trend.
"""

from __future__ import annotations

import re
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook

TAGS = ("BRAND", "GENERIC", "BRANDED GENERIC", "OTHER")
OA_CODES = frozenset({"M15", "M16", "M17", "M18", "M19"})
RA_CODES = frozenset({"M04"})
GRAND_TOTAL = "Grand Total"
OUTPUT_COLUMNS = [
    "disease_area",
    "manufacturer",
    "product",
    "brand_generic_tag",
    "icd10_code",
    "icd10_label",
    "patient_visits",
]

_ICD_HEADER = re.compile(r"^(M\d{2}) - (.+)\nPatient Visits$")


class ReferenceTableError(ValueError):
    """The workbook has no reference table, or it is not in the expected shape."""


def _trim(row: tuple) -> list:
    cells = list(row)
    while cells and cells[-1] is None:
        cells.pop()
    return cells


def _is_reference_header(row: tuple) -> bool:
    cells = _trim(row)
    if len(cells) < 4 or cells[:3] != ["Manufacturer", "Product Sum", "Brand/Generic"]:
        return False
    return all(isinstance(h, str) and _ICD_HEADER.match(h.strip()) for h in cells[3:])


def _disease_area(sheet_name: str, codes: list[str]) -> str:
    if set(codes) <= OA_CODES:
        from_codes = "OA"
    elif set(codes) <= RA_CODES:
        from_codes = "RA"
    else:
        raise ReferenceTableError(f"ICD-10 columns mix or do not belong to OA/RA: {codes}")
    if sheet_name.startswith("M15_19_OA"):
        from_sheet = "OA"
    elif sheet_name.startswith("M04_RA"):
        from_sheet = "RA"
    else:
        from_sheet = None
    if from_sheet is not None and from_sheet != from_codes:
        raise ReferenceTableError(
            f"Sheet {sheet_name!r} implies {from_sheet} but its ICD-10 columns imply {from_codes}"
        )
    return from_codes


def _read_reference_sheet(path: Path) -> tuple[str, list[tuple]]:
    workbook = load_workbook(path, read_only=True)
    try:
        matches = []
        for ws in workbook.worksheets:
            first = next(ws.iter_rows(min_row=1, max_row=1, values_only=True), ())
            if _is_reference_header(first):
                matches.append(ws)
        if len(matches) != 1:
            raise ReferenceTableError(
                f"Expected exactly one reference-table sheet, found {len(matches)}"
            )
        ws = matches[0]
        return ws.title, list(ws.iter_rows(values_only=True))
    finally:
        workbook.close()


def parse_reference_table(path: str | Path) -> pd.DataFrame:
    """Parse a Branded Generic workbook into a tidy long-format table.

    One row per (manufacturer, product, tag, ICD-10 code). A blank ICD-10 cell is emitted
    as 0 rather than dropped: in this table a row's existence is the information (it is the
    product master), so a product must not vanish because it has no visits under one code.

    The sheet is found by the shape of its header, never by name — the RA reference sheet
    carries the same name as the RA pivot sheet.

    Patient Visits is a distinct count, so rows overlap: each row is at most the Grand Total
    and the rows sum to at least the Grand Total. Equality is not expected and not checked.
    """
    sheet_name, rows = _read_reference_sheet(Path(path))
    header = _trim(rows[0])
    icd = [_ICD_HEADER.match(h.strip()).groups() for h in header[3:]]
    codes = [code for code, _ in icd]
    disease_area = _disease_area(sheet_name, codes)
    width = len(icd)

    detail = []
    grand_total = None
    for row_no, row in enumerate(rows[1:], start=2):
        cells = list(row) + [None] * (3 + width - len(row))
        if all(c is None for c in cells):
            continue
        if cells[0] == GRAND_TOTAL:
            if grand_total is not None:
                raise ReferenceTableError(f"Row {row_no}: second Grand Total row")
            grand_total = [v or 0 for v in cells[3 : 3 + width]]
            continue
        manufacturer, product, tag = cells[0], cells[1], cells[2]
        if not (isinstance(manufacturer, str) and isinstance(product, str) and manufacturer):
            raise ReferenceTableError(f"Row {row_no}: missing manufacturer or product")
        if tag not in TAGS:
            raise ReferenceTableError(f"Row {row_no}: unrecognized Brand/Generic tag {tag!r}")
        visits = []
        for value in cells[3 : 3 + width]:
            if value is None:
                visits.append(0)
            elif isinstance(value, int) and not isinstance(value, bool) and value >= 0:
                visits.append(value)
            else:
                raise ReferenceTableError(f"Row {row_no}: {value!r} is not a non-negative integer")
        detail.append((manufacturer, product, tag, visits))

    if grand_total is None:
        raise ReferenceTableError("No Grand Total row found")
    if not detail:
        raise ReferenceTableError("Reference table has no data rows")

    keys = [(m, p, t) for m, p, t, _ in detail]
    if len(keys) != len(set(keys)):
        raise ReferenceTableError("Duplicate (manufacturer, product, tag) rows")

    for col, (code, _) in enumerate(icd):
        column = [v[col] for *_, v in detail]
        if max(column) > grand_total[col]:
            raise ReferenceTableError(f"{code}: a row exceeds the Grand Total")
        if sum(column) < grand_total[col]:
            raise ReferenceTableError(f"{code}: rows sum to less than the Grand Total")

    records = [
        (disease_area, m, p, t, code, label, v[col])
        for m, p, t, v in detail
        for col, (code, label) in enumerate(icd)
    ]
    return pd.DataFrame(records, columns=OUTPUT_COLUMNS)
