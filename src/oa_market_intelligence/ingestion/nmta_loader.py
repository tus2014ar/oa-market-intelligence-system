"""Parser for the IQVIA NMTA nested-pivot patient-visit extracts (OA and RA).

Field-level structure is documented in docs/data_dictionary.md §2. The sheet is a
collapsed Excel PivotTable, not a flat table: Month, Manufacturer and Product rows
are interleaved in column A, and Month/Manufacturer rows hold subtotals.
"""

from __future__ import annotations

from datetime import datetime
from pathlib import Path

import numpy as np
import pandas as pd
from openpyxl import load_workbook

# The pivot marks its row hierarchy with the indent of the column-A cell.
MONTH_INDENT = 1
MANUFACTURER_INDENT = 3
PRODUCT_INDENT = 5

GRAND_TOTAL = "Grand Total"
VISITS_LABEL = "Patient Visits"

OUTPUT_COLUMNS = [
    "month",
    "disease_area",
    "manufacturer",
    "product",
    "specialty",
    "age_band",
    "gender",
    "patient_visits",
]


class PivotStructureError(ValueError):
    """The extract does not have the structure this parser was built against."""


def parse_header(header: str) -> tuple[str, str, str, str | None]:
    """Split a value-column header into (specialty, age_band, gender, icd10_label).

    OA headers have 4 newline-separated segments, RA headers have 5 (RA embeds the
    ICD-10 label). icd10_label is None for OA.
    """
    parts = [p.strip() for p in str(header).split("\n")]
    if len(parts) not in (4, 5) or parts[-1] != VISITS_LABEL:
        raise PivotStructureError(f"Unrecognized value-column header: {header!r}")
    icd = parts[3] if len(parts) == 5 else None
    return parts[0], parts[1], parts[2], icd


def detect_disease_area(sheet_name: str, icd_labels: list[str | None]) -> str:
    """Return "OA" or "RA" from header content, cross-checked against the sheet name.

    The file name is deliberately never consulted.
    """
    has_icd = [label is not None for label in icd_labels]
    if all(has_icd):
        from_headers = "RA"
    elif not any(has_icd):
        from_headers = "OA"
    else:
        raise PivotStructureError("Value-column headers mix 4- and 5-segment formats")

    if sheet_name.startswith("M15_19_OA"):
        from_sheet = "OA"
    elif sheet_name.startswith("M04_RA"):
        from_sheet = "RA"
    else:
        from_sheet = None
    if from_sheet is not None and from_sheet != from_headers:
        raise PivotStructureError(
            f"Sheet {sheet_name!r} implies {from_sheet} but its headers imply {from_headers}"
        )
    return from_headers


def _find_pivot_sheet(workbook):
    matches = []
    for ws in workbook.worksheets:
        first = next(ws.iter_rows(min_row=1, max_row=1, values_only=True), ())
        if first and first[0] == "Month" and any(
            isinstance(h, str) and h.endswith(VISITS_LABEL) for h in first[1:]
        ):
            matches.append(ws)
    if len(matches) != 1:
        raise PivotStructureError(
            f"Expected exactly one patient-visit pivot sheet, found {len(matches)}"
        )
    return matches[0]


def _read_rows(path: Path) -> tuple[str, list[tuple]]:
    workbook = load_workbook(path, read_only=True)
    try:
        ws = _find_pivot_sheet(workbook)
        rows = [
            (row[0].value, row[0].alignment.indent, [c.value for c in row[1:]])
            for row in ws.iter_rows()
        ]
        return ws.title, rows
    finally:
        workbook.close()


def _to_vector(values: list, width: int, row_no: int) -> np.ndarray:
    if len(values) != width:
        raise PivotStructureError(f"Row {row_no}: expected {width} value cells, got {len(values)}")
    try:
        return np.array([0 if v is None else v for v in values], dtype=np.int64)
    except (TypeError, ValueError) as exc:
        raise PivotStructureError(f"Row {row_no}: non-numeric visit count") from exc


def parse_pivot_sheet(path: str | Path) -> pd.DataFrame:
    """Parse an NMTA patient-visit pivot workbook into a tidy long-format table.

    One output row per (month, manufacturer, product, specialty, age band, gender)
    with a non-blank visit count. Blank pivot cells mean zero visits and are dropped.

    Patient Visits is a distinct count at every pivot level, so a visit involving two
    products of one manufacturer appears in both product rows but once in the
    manufacturer row. Product rows may therefore sum to more than their manufacturer
    subtotal; the parser checks the inequalities that must hold, not equality.
    """
    sheet_name, rows = _read_rows(Path(path))
    if not rows:
        raise PivotStructureError("Pivot sheet is empty")

    header_cells = rows[0][2]
    if any(h is None for h in header_cells):
        raise PivotStructureError("Header row has blank value-column headers")
    parsed_headers = [parse_header(h) for h in header_cells]
    disease_area = detect_disease_area(sheet_name, [h[3] for h in parsed_headers])
    width = len(parsed_headers)

    product_rows: list[tuple[datetime, str, str, np.ndarray]] = []
    month = manufacturer = None
    month_vec = mfr_vec = None
    mfr_product_sum = month_mfr_sum = None
    all_months_sum = np.zeros(width, dtype=np.int64)
    grand_total = None

    def close_manufacturer(row_no: int) -> None:
        nonlocal mfr_vec, mfr_product_sum, month_mfr_sum
        if mfr_vec is None:
            return
        if (mfr_vec > mfr_product_sum).any():
            raise PivotStructureError(
                f"Row {row_no}: manufacturer {manufacturer!r} subtotal exceeds the sum of "
                f"its product rows in {month:%Y-%m}"
            )
        month_mfr_sum += mfr_vec
        mfr_vec = None

    def close_month(row_no: int) -> None:
        nonlocal month_vec, all_months_sum
        close_manufacturer(row_no)
        if month_vec is None:
            return
        if (month_vec > month_mfr_sum).any():
            raise PivotStructureError(
                f"Row {row_no}: month total for {month:%Y-%m} exceeds the sum of its "
                "manufacturer rows"
            )
        all_months_sum += month_vec
        month_vec = None

    for offset, (label, indent, values) in enumerate(rows[1:]):
        row_no = offset + 2
        if indent == MONTH_INDENT and isinstance(label, datetime):
            close_month(row_no)
            month, month_vec = label, _to_vector(values, width, row_no)
            manufacturer, month_mfr_sum = None, np.zeros(width, dtype=np.int64)
        elif indent == MONTH_INDENT and label == GRAND_TOTAL:
            close_month(row_no)
            grand_total = _to_vector(values, width, row_no)
        elif indent == MANUFACTURER_INDENT and isinstance(label, str) and month_vec is not None:
            close_manufacturer(row_no)
            manufacturer, mfr_vec = label, _to_vector(values, width, row_no)
            mfr_product_sum = np.zeros(width, dtype=np.int64)
        elif indent == PRODUCT_INDENT and isinstance(label, str) and mfr_vec is not None:
            vec = _to_vector(values, width, row_no)
            if (vec > mfr_vec).any():
                raise PivotStructureError(
                    f"Row {row_no}: product {label!r} exceeds its manufacturer subtotal"
                )
            mfr_product_sum += vec
            product_rows.append((month, manufacturer, label, vec))
        else:
            raise PivotStructureError(
                f"Row {row_no}: unexpected row (label={label!r}, indent={indent})"
            )

    close_month(len(rows) + 1)
    if grand_total is None:
        raise PivotStructureError("No Grand Total row found")
    if not np.array_equal(grand_total, all_months_sum):
        raise PivotStructureError("Grand Total does not equal the sum of the month rows")
    if not product_rows:
        raise PivotStructureError("No product rows found")

    matrix = np.vstack([vec for *_, vec in product_rows])
    row_idx, col_idx = np.nonzero(matrix)
    months = np.array([r[0] for r in product_rows], dtype="datetime64[ns]")
    manufacturers = np.array([r[1] for r in product_rows], dtype=object)
    products = np.array([r[2] for r in product_rows], dtype=object)
    specialties = np.array([h[0] for h in parsed_headers], dtype=object)
    age_bands = np.array([h[1] for h in parsed_headers], dtype=object)
    genders = np.array([h[2] for h in parsed_headers], dtype=object)

    return pd.DataFrame(
        {
            "month": months[row_idx],
            "disease_area": disease_area,
            "manufacturer": manufacturers[row_idx],
            "product": products[row_idx],
            "specialty": specialties[col_idx],
            "age_band": age_bands[col_idx],
            "gender": genders[col_idx],
            "patient_visits": matrix[row_idx, col_idx],
        },
        columns=OUTPUT_COLUMNS,
    )
