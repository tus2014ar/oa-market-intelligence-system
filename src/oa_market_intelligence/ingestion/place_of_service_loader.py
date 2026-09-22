"""Loader for the market-level Place-of-Service-by-month summary sheet.

Every NMTA workbook carries this as a separate flat table on its secondary sheet
(docs/data_dictionary.md §3). It is not tied to any product, so it has its own grain:
one row per (month, disease area, setting).
"""

from __future__ import annotations

from datetime import datetime
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook

SETTINGS = ("HOSPITAL", "OFFICE", "OTHER", "TELEHEALTH")
OUTPUT_COLUMNS = ["month", "disease_area", "place_of_service", "patient_visits"]

# The header row is at row 29 in every extract seen so far, but scan for it rather than
# hard-code that; the cap keeps the scan from reading the 14,000-row pivot sheet.
_HEADER_SCAN_ROWS = 100


class PlaceOfServiceError(ValueError):
    """The workbook's Place-of-Service sheet is missing or not in the expected shape."""


def _find_header(rows: list[tuple]) -> int | None:
    for index, row in enumerate(rows):
        if not row or row[0] != "Month":
            continue
        headers = [h for h in row[1:] if h is not None]
        if headers and all(h in SETTINGS for h in headers):
            return index
    return None


def _disease_area(sheet_name: str) -> str:
    if sheet_name.startswith("M15_19_OA"):
        return "OA"
    if sheet_name.startswith("M04_RA"):
        return "RA"
    raise PlaceOfServiceError(f"Cannot tell OA from RA using sheet name {sheet_name!r}")


def _parse_month(label, row_no: int) -> datetime:
    if isinstance(label, datetime):
        return label
    try:
        return datetime.strptime(str(label).strip(), "%b %Y")
    except ValueError:
        raise PlaceOfServiceError(
            f"Row {row_no}: month label {label!r} is not 'Mon YYYY'. Truncated labels such "
            "as 'Sep...' (the Branded Generic workbooks) carry no year and are not loaded."
        ) from None


def parse_place_of_service(path: str | Path) -> pd.DataFrame:
    """Parse a workbook's Place-of-Service sheet into a tidy long-format table.

    Blank cells mean zero visits and are dropped. A setting with no column at all
    (the RA sheet has no HOSPITAL or TELEHEALTH) is simply absent, not an error.
    """
    workbook = load_workbook(Path(path), read_only=True)
    try:
        candidates = []
        for ws in workbook.worksheets:
            rows = list(ws.iter_rows(max_row=_HEADER_SCAN_ROWS, values_only=True))
            header_index = _find_header(rows)
            if header_index is not None:
                candidates.append(ws)
        if len(candidates) != 1:
            raise PlaceOfServiceError(
                f"Expected exactly one Place-of-Service sheet, found {len(candidates)}"
            )
        ws = candidates[0]
        all_rows = list(ws.iter_rows(values_only=True))
        sheet_name = ws.title
    finally:
        workbook.close()

    disease_area = _disease_area(sheet_name)
    header_index = _find_header(all_rows)
    headers = all_rows[header_index]
    settings = [(col, h) for col, h in enumerate(headers) if col > 0 and h is not None]

    records = []
    seen_months = set()
    for offset, row in enumerate(all_rows[header_index + 1 :]):
        if all(v is None for v in row):
            continue
        row_no = header_index + offset + 2
        month = _parse_month(row[0], row_no)
        if month in seen_months:
            raise PlaceOfServiceError(f"Row {row_no}: month {month:%Y-%m} appears twice")
        seen_months.add(month)
        for col, setting in settings:
            value = row[col]
            if value is None:
                continue
            if not isinstance(value, int) or isinstance(value, bool) or value < 0:
                raise PlaceOfServiceError(
                    f"Row {row_no}: {setting} value {value!r} is not a non-negative integer"
                )
            records.append((month, disease_area, setting, value))

    if not seen_months:
        raise PlaceOfServiceError("Place-of-Service sheet has no data rows")

    frame = pd.DataFrame(records, columns=OUTPUT_COLUMNS)
    frame["month"] = pd.to_datetime(frame["month"]).astype("datetime64[ns]")
    return frame
