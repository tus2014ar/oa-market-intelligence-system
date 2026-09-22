"""Tests for the Branded Generic reference-table loader, run against the real extracts."""

from pathlib import Path

import pandas as pd
import pytest
from openpyxl import load_workbook

from oa_market_intelligence.ingestion.nmta_loader import PivotStructureError, parse_pivot_sheet
from oa_market_intelligence.ingestion.reference_loader import (
    OUTPUT_COLUMNS,
    ReferenceTableError,
    parse_reference_table,
)

ROOT = Path(__file__).resolve().parent.parent
RAW = ROOT / "data" / "raw"
OA_REF = RAW / "Branded Generic - OA.xlsx"
RA_REF = RAW / "Branded Generic - RA.xlsx"
OA_PIVOT = RAW / "Team1_M15_19_OA.xlsx"
RA_PIVOT = RAW / "Team1_M04_RA.xlsx"
OA_REF_SHEET = "M15_19_OA_BRANDED_GENERIC"
ENTITY = ["manufacturer", "product", "brand_generic_tag"]


@pytest.fixture(scope="module")
def oa_ref():
    return parse_reference_table(OA_REF)


@pytest.fixture(scope="module")
def ra_ref():
    return parse_reference_table(RA_REF)


@pytest.fixture(scope="module")
def taxonomy():
    return pd.read_csv(ROOT / "data" / "reference" / "product_taxonomy.csv")


def _corrupted_oa_copy(tmp_path, mutate):
    workbook = load_workbook(OA_REF)
    mutate(workbook[OA_REF_SHEET])
    out = tmp_path / "corrupted.xlsx"
    workbook.save(out)
    return out


# ---------- real OA reference ----------


def test_oa_shape_and_columns(oa_ref):
    assert list(oa_ref.columns) == OUTPUT_COLUMNS
    assert (oa_ref["disease_area"] == "OA").all()
    assert oa_ref["product"].nunique() == 145
    assert oa_ref["manufacturer"].nunique() == 168
    assert set(oa_ref["icd10_code"]) == {"M16", "M17"}
    assert len(oa_ref.drop_duplicates(ENTITY)) == 417
    assert len(oa_ref) == 417 * 2


def test_oa_rows_overlap_so_they_exceed_the_printed_grand_total(oa_ref):
    by_code = oa_ref.groupby("icd10_code")["patient_visits"].sum()
    assert by_code["M16"] == 567_927 > 525_661
    assert by_code["M17"] == 4_993_204 > 4_797_621


def test_blank_icd_cells_become_zero_not_dropped(oa_ref):
    first = oa_ref[(oa_ref["product"] == "ANJESO") & (oa_ref["manufacturer"] == "No Manufacturer")]
    assert first.set_index("icd10_code")["patient_visits"].to_dict() == {"M16": 0, "M17": 1}


def test_a_product_can_carry_more_than_one_tag(oa_ref):
    tags = oa_ref.drop_duplicates(ENTITY).groupby("product")["brand_generic_tag"].nunique()
    assert (tags > 1).sum() == 7
    acetaminophen = set(oa_ref[oa_ref["product"] == "ACETAMINOPHEN"]["brand_generic_tag"])
    assert acetaminophen == {"OTHER", "GENERIC"}


def test_tags_are_patent_status_not_drug_class(oa_ref):
    tag = oa_ref.drop_duplicates("product").set_index("product")["brand_generic_tag"]
    assert tag["KENALOG"] == "BRANDED GENERIC"
    assert tag["DEPO-MEDROL"] == "BRAND"
    assert tag["ZILRETTA"] == "BRANDED GENERIC"


def test_zilretta_is_split_across_two_manufacturer_labels(oa_ref):
    zilretta = oa_ref[oa_ref["product"] == "ZILRETTA"]
    assert set(zilretta["manufacturer"]) == {"No Manufacturer", "PACIRA PHARM"}


def test_oa_products_match_the_committed_taxonomy_exactly(oa_ref, taxonomy):
    in_taxonomy = set(taxonomy[taxonomy["disease_area"] == "OA"]["product_name"])
    assert set(oa_ref["product"]) == in_taxonomy


# ---------- real RA reference ----------


def test_ra_is_all_brand_biologics(ra_ref):
    assert (ra_ref["disease_area"] == "RA").all()
    assert set(ra_ref["icd10_code"]) == {"M04"}
    assert ra_ref["product"].nunique() == 15
    assert set(ra_ref["brand_generic_tag"]) == {"BRAND"}
    assert len(ra_ref) == 18
    assert ra_ref["patient_visits"].sum() == 1_287


def test_ra_products_match_the_committed_taxonomy_exactly(ra_ref, taxonomy):
    in_taxonomy = set(taxonomy[taxonomy["disease_area"] == "RA"]["product_name"])
    assert set(ra_ref["product"]) == in_taxonomy


# ---------- the RA sheet-name collision ----------


def test_ra_reference_and_ra_pivot_share_a_sheet_name_but_are_told_apart_by_shape():
    ref_sheets = load_workbook(RA_REF, read_only=True).sheetnames
    pivot_sheets = load_workbook(RA_PIVOT, read_only=True).sheetnames
    assert "M04_RA_PAT_VISIT" in ref_sheets and "M04_RA_PAT_VISIT" in pivot_sheets

    with pytest.raises(ReferenceTableError):
        parse_reference_table(RA_PIVOT)
    with pytest.raises(PivotStructureError):
        parse_pivot_sheet(RA_REF)


def test_refuses_the_oa_pivot_workbook():
    with pytest.raises(ReferenceTableError, match="reference-table sheet"):
        parse_reference_table(OA_PIVOT)


# ---------- malformed input must halt ----------


def test_rejects_an_unrecognized_tag(tmp_path):
    def mutate(ws):
        ws.cell(row=2, column=3).value = "WEIRD"

    with pytest.raises(ReferenceTableError, match="unrecognized Brand/Generic tag"):
        parse_reference_table(_corrupted_oa_copy(tmp_path, mutate))


def test_rejects_a_duplicated_row(tmp_path):
    def mutate(ws):
        for column in (1, 2, 3):
            ws.cell(row=3, column=column).value = ws.cell(row=2, column=column).value

    with pytest.raises(ReferenceTableError, match="Duplicate"):
        parse_reference_table(_corrupted_oa_copy(tmp_path, mutate))


def test_rejects_a_row_larger_than_the_grand_total(tmp_path):
    def mutate(ws):
        ws.cell(row=2, column=5).value = 100_000_000

    with pytest.raises(ReferenceTableError, match="exceeds the Grand Total"):
        parse_reference_table(_corrupted_oa_copy(tmp_path, mutate))


def test_rejects_rows_that_sum_to_less_than_the_grand_total(tmp_path):
    def mutate(ws):
        ws.cell(row=ws.max_row, column=5).value += 1_000_000

    with pytest.raises(ReferenceTableError, match="less than the Grand Total"):
        parse_reference_table(_corrupted_oa_copy(tmp_path, mutate))


def test_rejects_a_missing_grand_total(tmp_path):
    def mutate(ws):
        ws.delete_rows(ws.max_row)

    with pytest.raises(ReferenceTableError, match="No Grand Total"):
        parse_reference_table(_corrupted_oa_copy(tmp_path, mutate))


def test_rejects_a_negative_visit_count(tmp_path):
    def mutate(ws):
        ws.cell(row=2, column=5).value = -5

    with pytest.raises(ReferenceTableError, match="non-negative integer"):
        parse_reference_table(_corrupted_oa_copy(tmp_path, mutate))


def test_rejects_icd_columns_that_mix_oa_and_ra(tmp_path):
    def mutate(ws):
        ws.cell(row=1, column=4).value = "M04 - AUTOINFLAMMATORY SYNDROMES\nPatient Visits"

    with pytest.raises(ReferenceTableError, match="mix or do not belong"):
        parse_reference_table(_corrupted_oa_copy(tmp_path, mutate))


def test_rejects_a_sheet_name_that_contradicts_the_icd_columns(tmp_path):
    def mutate(ws):
        ws.title = "M04_RA_BRANDED_GENERIC"

    workbook = load_workbook(OA_REF)
    mutate(workbook[OA_REF_SHEET])
    out = tmp_path / "renamed.xlsx"
    workbook.save(out)
    with pytest.raises(ReferenceTableError, match="implies RA"):
        parse_reference_table(out)
