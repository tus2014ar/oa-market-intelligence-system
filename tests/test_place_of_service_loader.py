"""Tests for the Place-of-Service loader, run against the real committed extracts."""

import shutil
from pathlib import Path

import pytest
from openpyxl import load_workbook

from oa_market_intelligence.ingestion.place_of_service_loader import (
    OUTPUT_COLUMNS,
    PlaceOfServiceError,
    parse_place_of_service,
)

RAW = Path(__file__).resolve().parent.parent / "data" / "raw"
OA_FILE = RAW / "Team1_M15_19_OA.xlsx"
RA_FILE = RAW / "Team1_M04_RA.xlsx"
RA_POS_SHEET = "M04_RA_PAT_VISIT1"
HEADER_ROW = 29


@pytest.fixture(scope="module")
def oa_pos():
    return parse_place_of_service(OA_FILE)


@pytest.fixture(scope="module")
def ra_pos():
    return parse_place_of_service(RA_FILE)


def _corrupted_copy(tmp_path, mutate):
    workbook = load_workbook(RA_FILE)
    mutate(workbook)
    out = tmp_path / "corrupted.xlsx"
    workbook.save(out)
    return out


# ---------- real OA extract ----------


def test_oa_has_all_72_months_and_four_settings(oa_pos):
    assert list(oa_pos.columns) == OUTPUT_COLUMNS
    assert (oa_pos["disease_area"] == "OA").all()
    assert oa_pos["month"].nunique() == 72
    assert oa_pos["month"].min().strftime("%Y-%m") == "2019-08"
    assert oa_pos["month"].max().strftime("%Y-%m") == "2025-07"
    assert set(oa_pos["place_of_service"]) == {"HOSPITAL", "OFFICE", "OTHER", "TELEHEALTH"}


def test_oa_total_matches_the_documented_market_total(oa_pos):
    assert oa_pos["patient_visits"].sum() == 7_189_004


def test_oa_shows_the_april_2020_office_to_telehealth_shock(oa_pos):
    april = oa_pos[oa_pos["month"] == "2020-04-01"].set_index("place_of_service")
    assert april.loc["OFFICE", "patient_visits"] == 39_487
    assert april.loc["TELEHEALTH", "patient_visits"] == 1_016


def test_month_column_matches_the_pivot_parser_dtype(oa_pos):
    assert str(oa_pos["month"].dtype) == "datetime64[ns]"


def test_keys_are_unique_and_counts_positive(oa_pos):
    assert not oa_pos.duplicated(["month", "disease_area", "place_of_service"]).any()
    assert (oa_pos["patient_visits"] > 0).all()


# ---------- real RA extract ----------


def test_ra_tolerates_missing_setting_columns(ra_pos):
    assert (ra_pos["disease_area"] == "RA").all()
    assert set(ra_pos["place_of_service"]) == {"OFFICE", "OTHER"}


def test_ra_months_with_no_visits_are_absent_not_zero_filled(ra_pos):
    assert ra_pos["month"].nunique() < 72


def test_disease_area_ignores_the_file_name(tmp_path):
    misleading = tmp_path / "Team1_M15_19_OA.xlsx"
    shutil.copy(RA_FILE, misleading)
    assert (parse_place_of_service(misleading)["disease_area"] == "RA").all()


# ---------- inputs that must be refused ----------


@pytest.mark.parametrize("name", ["Branded Generic - OA.xlsx", "Branded Generic - RA.xlsx"])
def test_refuses_reference_workbooks_with_year_less_month_labels(name):
    with pytest.raises(PlaceOfServiceError, match="carry no year"):
        parse_place_of_service(RAW / name)


def test_rejects_a_negative_visit_count(tmp_path):
    def mutate(wb):
        wb[RA_POS_SHEET].cell(row=HEADER_ROW + 1, column=2).value = -1

    with pytest.raises(PlaceOfServiceError, match="non-negative integer"):
        parse_place_of_service(_corrupted_copy(tmp_path, mutate))


def test_rejects_a_duplicated_month(tmp_path):
    def mutate(wb):
        ws = wb[RA_POS_SHEET]
        ws.cell(row=HEADER_ROW + 2, column=1).value = ws.cell(row=HEADER_ROW + 1, column=1).value

    with pytest.raises(PlaceOfServiceError, match="appears twice"):
        parse_place_of_service(_corrupted_copy(tmp_path, mutate))


def test_rejects_an_unknown_setting_header(tmp_path):
    def mutate(wb):
        wb[RA_POS_SHEET].cell(row=HEADER_ROW, column=2).value = "PHARMACY"

    with pytest.raises(PlaceOfServiceError, match="exactly one Place-of-Service sheet"):
        parse_place_of_service(_corrupted_copy(tmp_path, mutate))


def test_rejects_an_unrecognizable_sheet_name(tmp_path):
    def mutate(wb):
        wb[RA_POS_SHEET].title = "Sheet2"

    with pytest.raises(PlaceOfServiceError, match="OA from RA"):
        parse_place_of_service(_corrupted_copy(tmp_path, mutate))


def test_rejects_a_workbook_with_no_place_of_service_sheet(tmp_path):
    def mutate(wb):
        del wb[RA_POS_SHEET]

    with pytest.raises(PlaceOfServiceError, match="exactly one Place-of-Service sheet"):
        parse_place_of_service(_corrupted_copy(tmp_path, mutate))
