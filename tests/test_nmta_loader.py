"""Tests for the NMTA pivot parser, run against the real committed extracts."""

import shutil
from pathlib import Path

import pytest
from openpyxl import load_workbook
from openpyxl.styles import Alignment

from oa_market_intelligence.ingestion.nmta_loader import (
    OUTPUT_COLUMNS,
    PivotStructureError,
    detect_disease_area,
    parse_header,
    parse_pivot_sheet,
)

RAW = Path(__file__).resolve().parent.parent / "data" / "raw"
OA_FILE = RAW / "Team1_M15_19_OA.xlsx"
RA_FILE = RAW / "Team1_M04_RA.xlsx"
REFERENCE_FILE = RAW / "Branded Generic - RA.xlsx"
RA_SHEET = "M04_RA_PAT_VISIT"


@pytest.fixture(scope="module")
def oa():
    return parse_pivot_sheet(OA_FILE)


@pytest.fixture(scope="module")
def ra():
    return parse_pivot_sheet(RA_FILE)


# ---------- header parsing and disease-area detection ----------


def test_parse_header_oa_has_no_icd_label():
    header = "ADDICTION MEDICINE\n40 TO 59\nFEMALE\nPatient Visits"
    assert parse_header(header) == ("ADDICTION MEDICINE", "40 TO 59", "FEMALE", None)


def test_parse_header_ra_carries_icd_label():
    header = "ALLERGY\n00 TO 02\nMALE\nM04 - AUTOINFLAMMATORY SYNDROMES\nPatient Visits"
    assert parse_header(header) == (
        "ALLERGY",
        "00 TO 02",
        "MALE",
        "M04 - AUTOINFLAMMATORY SYNDROMES",
    )


@pytest.mark.parametrize(
    "bad",
    ["ONLY\nTWO", "A\nB\nC\nNot Visits", "A\nB\nC\nD\nE\nPatient Visits", "no newlines at all"],
)
def test_parse_header_rejects_unrecognized_formats(bad):
    with pytest.raises(PivotStructureError):
        parse_header(bad)


def test_detect_disease_area_from_headers():
    assert detect_disease_area("M15_19_OA_PAT_VISIT", [None, None]) == "OA"
    assert detect_disease_area("M04_RA_PAT_VISIT", ["M04", "M04"]) == "RA"


def test_detect_disease_area_falls_back_to_headers_for_unknown_sheet_name():
    assert detect_disease_area("Sheet1", ["M04"]) == "RA"


def test_detect_disease_area_rejects_sheet_header_mismatch():
    with pytest.raises(PivotStructureError):
        detect_disease_area("M15_19_OA_PAT_VISIT", ["M04"])


def test_detect_disease_area_rejects_mixed_headers():
    with pytest.raises(PivotStructureError):
        detect_disease_area("M04_RA_PAT_VISIT", ["M04", None])


# ---------- real RA extract ----------


def test_ra_matches_data_dictionary(ra):
    assert list(ra.columns) == OUTPUT_COLUMNS
    assert (ra["disease_area"] == "RA").all()
    assert ra["month"].nunique() == 72
    assert ra["month"].min().strftime("%Y-%m") == "2019-08"
    assert ra["month"].max().strftime("%Y-%m") == "2025-07"
    assert ra["specialty"].nunique() == 19
    assert ra["age_band"].nunique() == 9
    assert set(ra["gender"]) == {"FEMALE", "MALE"}


def test_ra_visit_counts_are_positive_and_keys_unique(ra):
    assert (ra["patient_visits"] > 0).all()
    key = ["month", "manufacturer", "product", "specialty", "age_band", "gender"]
    assert not ra.duplicated(key).any()


def test_ra_includes_the_miscoded_oncology_products(ra):
    assert {"OPDIVO", "YERVOY", "KEYTRUDA"} <= set(ra["product"])


def test_disease_area_ignores_the_file_name(tmp_path):
    misleading = tmp_path / "Team1_M15_19_OA.xlsx"
    shutil.copy(RA_FILE, misleading)
    assert (parse_pivot_sheet(misleading)["disease_area"] == "RA").all()


# ---------- real OA extract ----------


def test_oa_matches_data_dictionary(oa):
    assert list(oa.columns) == OUTPUT_COLUMNS
    assert (oa["disease_area"] == "OA").all()
    assert oa["month"].nunique() == 72
    assert oa["month"].min().strftime("%Y-%m") == "2019-08"
    assert oa["month"].max().strftime("%Y-%m") == "2025-07"
    assert oa["specialty"].nunique() == 50
    assert oa["age_band"].nunique() == 10
    assert set(oa["gender"]) == {"FEMALE", "MALE", "UNSPECIFIED"}
    assert len(oa[["specialty", "age_band", "gender"]].drop_duplicates()) == 647


def test_oa_visit_counts_are_positive_and_keys_unique(oa):
    assert (oa["patient_visits"] > 0).all()
    key = ["month", "manufacturer", "product", "specialty", "age_band", "gender"]
    assert not oa.duplicated(key).any()


def test_oa_zilretta_is_split_across_two_manufacturer_labels(oa):
    zilretta = oa[oa["product"] == "ZILRETTA"]
    assert set(zilretta["manufacturer"]) == {"No Manufacturer", "PACIRA PHARM"}


# ---------- malformed input must halt, not parse silently ----------


def _corrupted_copy(tmp_path, mutate):
    workbook = load_workbook(RA_FILE)
    mutate(workbook[RA_SHEET])
    out = tmp_path / "corrupted.xlsx"
    workbook.save(out)
    return out


def _first_row_with_indent(ws, indent):
    return next(row[0].row for row in ws.iter_rows() if row[0].alignment.indent == indent)


def test_rejects_unexpected_row_indent(tmp_path):
    def mutate(ws):
        row = _first_row_with_indent(ws, 3)
        ws.cell(row=row, column=1).alignment = Alignment(indent=2)

    with pytest.raises(PivotStructureError, match="unexpected row"):
        parse_pivot_sheet(_corrupted_copy(tmp_path, mutate))


def test_rejects_product_exceeding_manufacturer_subtotal(tmp_path):
    def mutate(ws):
        row = _first_row_with_indent(ws, 5)
        cell = next(c for c in ws[row][1:] if c.value is not None)
        cell.value += 1000

    with pytest.raises(PivotStructureError, match="exceeds its manufacturer subtotal"):
        parse_pivot_sheet(_corrupted_copy(tmp_path, mutate))


def test_rejects_grand_total_that_does_not_reconcile(tmp_path):
    def mutate(ws):
        cell = next(c for c in ws[ws.max_row][1:] if c.value is not None)
        cell.value += 1

    with pytest.raises(PivotStructureError, match="Grand Total"):
        parse_pivot_sheet(_corrupted_copy(tmp_path, mutate))


def test_rejects_corrupted_value_column_header(tmp_path):
    def mutate(ws):
        ws["B1"] = "NOT A VALID HEADER"

    with pytest.raises(PivotStructureError, match="Unrecognized value-column header"):
        parse_pivot_sheet(_corrupted_copy(tmp_path, mutate))


def test_rejects_a_file_with_no_pivot_sheet():
    with pytest.raises(PivotStructureError, match="patient-visit pivot sheet"):
        parse_pivot_sheet(REFERENCE_FILE)
